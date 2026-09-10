#!/usr/bin/env python3
"""Merge Pepaws Porch submissions into the backend workbook.

Getting the input file:

  1. Open the game with #assessor on the end of its URL, in a NEW tab.
     (Typing the hash into a tab that already has the game open works too now,
     but a fresh tab is the reliable habit.)
  2. Click "Export raw JSON". Confirm the download.
  3. Point this script at the file you just downloaded:

         python3 scripts/sync_responses.py ~/Downloads/pepaws-porch-submissions.json

It also accepts a single submission document, or a directory containing either.

Rows are keyed on participant id, so re-running with the same export is safe —
an existing participant's rows are replaced, not duplicated, and the four
assessor columns are never overwritten. The workbook is created on first run.

If you only want a spreadsheet to read, you do not need this script at all:
"Export workbook (.xlsx)" on the same page gives you one directly.

Sheets
    Submissions   one row per participant, plus the assessor's own columns
    Free text     one row per written answer, with the rubric signals it hit
    Choices       every choice made, with the meter movement it caused
    Rubric        the signal set the game actually scores against
"""

import argparse
import json
import os
import sys
from datetime import datetime, timezone

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required:  pip install openpyxl")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_WORKBOOK = os.path.join(ROOT, "data", "pepaws-porch-responses.xlsx")
RUBRIC_JSON = os.path.join(ROOT, "data", "rubric.json")

HEAD_FILL = PatternFill("solid", fgColor="E7DFCF")
HEAD_FONT = Font(bold=True, color="31302B")

SUBMISSION_COLS = [
    ("participant_id", 22), ("name", 22), ("email", 28), ("submitted_at", 20),
    ("duration_min", 12), ("rapport", 9), ("autonomy", 10), ("safety", 9),
    ("provisional_band", 20), ("hard_stop", 10), ("declined", 10),
    ("watch_flags", 30), ("provisional_notes", 70),
    ("assessor", 16), ("assessor_outcome", 20), ("assessor_notes", 50), ("reviewed_at", 18),
]
FREETEXT_COLS = [
    ("participant_id", 22), ("name", 20), ("level", 8), ("day", 6),
    ("competency", 40), ("prompt", 44), ("answer", 90), ("characters", 11),
    ("seconds", 9), ("signals_matched", 32), ("signals_missed", 32), ("watch_flags", 24),
]
CHOICE_COLS = [
    ("participant_id", 22), ("day", 6), ("time", 10), ("scene", 22), ("level", 7),
    ("prompt", 42), ("chose", 58), ("tag", 15),
    ("rapport_delta", 13), ("autonomy_delta", 14), ("safety_delta", 12),
    ("timed_out", 10), ("seconds", 9),
]
# Columns an assessor fills in by hand — never overwritten by a re-sync.
ASSESSOR_COLS = ["assessor", "assessor_outcome", "assessor_notes", "reviewed_at"]


def style_header(ws, cols):
    for i, (name, width) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="top")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"


def ensure_sheet(wb, title, cols):
    if title in wb.sheetnames:
        return wb[title]
    ws = wb.create_sheet(title)
    style_header(ws, cols)
    return ws


def read_rows(ws, cols):
    """Existing data rows as dicts, keyed in sheet order."""
    names = [c[0] for c in cols]
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None or all(v is None for v in row):
            continue
        out.append({names[i]: (row[i] if i < len(row) else None) for i in range(len(names))})
    return out


def write_rows(ws, cols, rows):
    ws.delete_rows(2, max(ws.max_row - 1, 1))
    for r, row in enumerate(rows, start=2):
        for c, (name, _w) in enumerate(cols, start=1):
            ws.cell(row=r, column=c, value=row.get(name))


def unwrap(doc):
    """Peel any wrapper an export tool put around the submission body."""
    for _ in range(4):
        if not isinstance(doc, dict) or doc.get("pid"):
            break
        for key in ("data", "document", "fields", "value"):
            inner = doc.get(key)
            if isinstance(inner, dict):
                doc = inner
                break
        else:
            break
    return doc


def load_documents(path):
    """Accepts, in any combination:
         - the assessor view's "Export raw JSON" file (a list of submissions)
         - a single submission document
         - a directory of either
    """
    docs, skipped = [], []
    if os.path.isfile(path):
        candidates = [path]
    else:
        candidates = []
        for base, _dirs, files in os.walk(path):
            for f in sorted(files):
                if f.lower().endswith(".json"):
                    candidates.append(os.path.join(base, f))
        if not candidates:
            skipped.append(f"{path}: no .json files found in this directory")

    for f in candidates:
        name = os.path.basename(f)
        try:
            with open(f, encoding="utf-8") as fh:
                payload = json.load(fh)
        except (OSError, ValueError) as e:
            skipped.append(f"{name}: {e}")
            continue

        # A list of submissions, or an object wrapping one.
        if isinstance(payload, dict):
            for key in ("documents", "submissions", "records", "docs", "items"):
                if isinstance(payload.get(key), list):
                    payload = payload[key]
                    break
        items = payload if isinstance(payload, list) else [payload]

        found = 0
        for item in items:
            doc = unwrap(item)
            if isinstance(doc, dict) and doc.get("pid"):
                docs.append(doc)
                found += 1
        if not found:
            skipped.append(f"{name}: no submissions found (no participant id in it)")
    return docs, skipped


def joinlist(v):
    return ", ".join(str(x) for x in v) if isinstance(v, list) else (v or "")


def submission_row(doc):
    return {
        "participant_id": doc.get("pid"),
        "name": doc.get("name"),
        "email": doc.get("email"),
        "submitted_at": (doc.get("submittedAt") or "").replace("T", " ")[:19],
        "duration_min": doc.get("durationMin"),
        "rapport": doc.get("rapport"),
        "autonomy": doc.get("autonomy"),
        "safety": doc.get("safety"),
        "provisional_band": doc.get("band"),
        "hard_stop": "yes" if doc.get("hardStop") else "",
        "declined": "yes" if doc.get("declined") else "",
        "watch_flags": joinlist(doc.get("watchFlags")),
        "provisional_notes": " | ".join(doc.get("reasons") or []),
        "assessor": None, "assessor_outcome": None, "assessor_notes": None, "reviewed_at": None,
    }


def freetext_rows(doc):
    rows = []
    for r in doc.get("responses") or []:
        rows.append({
            "participant_id": doc.get("pid"), "name": doc.get("name"),
            "level": r.get("id"), "day": r.get("day"),
            "competency": r.get("competency"), "prompt": r.get("prompt"),
            "answer": r.get("answer"), "characters": r.get("chars"), "seconds": r.get("seconds"),
            "signals_matched": joinlist(r.get("matched")),
            "signals_missed": joinlist(r.get("missed")),
            "watch_flags": joinlist(r.get("watch")),
        })
    return rows


def choice_rows(doc):
    rows = []
    for c in doc.get("choices") or []:
        rows.append({
            "participant_id": doc.get("pid"), "day": c.get("day"), "time": c.get("time"),
            "scene": c.get("title"), "level": c.get("level"), "prompt": c.get("prompt"),
            "chose": c.get("chosen"), "tag": c.get("tag"),
            "rapport_delta": c.get("dr"), "autonomy_delta": c.get("da"), "safety_delta": c.get("ds"),
            "timed_out": "yes" if c.get("timedOut") else "", "seconds": c.get("seconds"),
        })
    return rows


def rebuild_rubric_sheet(wb):
    if "Rubric" in wb.sheetnames:
        del wb["Rubric"]
    ws = wb.create_sheet("Rubric")
    cols = [("level", 8), ("competency", 46), ("kind", 9), ("signal_id", 24),
            ("what_it_looks_like", 54), ("pattern", 60)]
    style_header(ws, cols)
    if not os.path.exists(RUBRIC_JSON):
        ws.cell(row=2, column=1, value="run: node scripts/export-rubric.js")
        return ws
    with open(RUBRIC_JSON, encoding="utf-8") as fh:
        data = json.load(fh)
    r = 2
    for level, spec in data.get("rubric", {}).items():
        for kind in ("signals", "watch"):
            for s in spec.get(kind, []):
                for c, v in enumerate([level, spec.get("competency"),
                                       "signal" if kind == "signals" else "watch",
                                       s["id"], s["label"], s["pattern"]], start=1):
                    ws.cell(row=r, column=c, value=v)
                r += 1
    r += 1
    ws.cell(row=r, column=1, value="Bands").font = HEAD_FONT
    r += 1
    for b in data.get("bands", []):
        ws.cell(row=r, column=1, value=b["band"])
        ws.cell(row=r, column=2, value=b["meaning"])
        r += 1
    return ws


def main():
    ap = argparse.ArgumentParser(description="Merge Pepaws Porch submissions into the backend workbook.")
    ap.add_argument("json_dir", nargs="?", metavar="JSON",
                    help='the "Export raw JSON" file from the game\'s #assessor page '
                         "(or a single submission document, or a directory of either)")
    ap.add_argument("--workbook", default=DEFAULT_WORKBOOK, help="workbook to create or update")
    args = ap.parse_args()

    if os.path.exists(args.workbook):
        wb = load_workbook(args.workbook)
    else:
        wb = Workbook()
        wb.remove(wb.active)
        os.makedirs(os.path.dirname(args.workbook), exist_ok=True)

    subs = ensure_sheet(wb, "Submissions", SUBMISSION_COLS)
    ft = ensure_sheet(wb, "Free text", FREETEXT_COLS)
    ch = ensure_sheet(wb, "Choices", CHOICE_COLS)
    rebuild_rubric_sheet(wb)
    wb._sheets = [wb[n] for n in ["Submissions", "Free text", "Choices", "Rubric"]]

    existing_subs = read_rows(subs, SUBMISSION_COLS)
    existing_ft = read_rows(ft, FREETEXT_COLS)
    existing_ch = read_rows(ch, CHOICE_COLS)

    docs, skipped = ([], [])
    if args.json_dir:
        if not os.path.exists(args.json_dir):
            sys.exit(f"no such path: {args.json_dir}")
        docs, skipped = load_documents(args.json_dir)

    incoming = {d["pid"]: d for d in docs}   # last document per pid wins
    # Keep the assessor's hand-written columns across a re-sync.
    kept_assessor = {r["participant_id"]: {k: r.get(k) for k in ASSESSOR_COLS}
                     for r in existing_subs if r.get("participant_id") in incoming}

    new_subs = [r for r in existing_subs if r.get("participant_id") not in incoming]
    new_ft = [r for r in existing_ft if r.get("participant_id") not in incoming]
    new_ch = [r for r in existing_ch if r.get("participant_id") not in incoming]

    for pid, doc in incoming.items():
        row = submission_row(doc)
        row.update({k: v for k, v in kept_assessor.get(pid, {}).items() if v})
        new_subs.append(row)
        new_ft.extend(freetext_rows(doc))
        new_ch.extend(choice_rows(doc))

    new_subs.sort(key=lambda r: str(r.get("submitted_at") or ""), reverse=True)

    write_rows(subs, SUBMISSION_COLS, new_subs)
    write_rows(ft, FREETEXT_COLS, new_ft)
    write_rows(ch, CHOICE_COLS, new_ch)
    for ws, wrapcol in ((ft, "G"), (subs, "M")):
        for cell in ws[wrapcol][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(args.workbook)

    stamp = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    print(f"{args.workbook}")
    print(f"  {len(incoming)} submission(s) merged, {len(new_subs)} participant(s) total  [{stamp}]")
    print(f"  free-text rows: {len(new_ft)}   choice rows: {len(new_ch)}")
    for s in skipped:
        print(f"  skipped {s}")
    if args.json_dir and not incoming:
        print()
        print("  Nothing was merged. That input held no submissions.")
        print("  Get the right file: open the game with #assessor on the end of its URL,")
        print("  then click 'Export raw JSON' and pass the downloaded file to this script.")


if __name__ == "__main__":
    main()
